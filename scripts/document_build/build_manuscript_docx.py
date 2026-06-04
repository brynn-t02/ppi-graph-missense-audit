from pathlib import Path
import sys
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "manuscript_english_working_draft_v1.md"
OUTPUT = ROOT / "manuscript_english_working_draft_v1.docx"


def set_font(run, name="Calibri", size=11, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, *, after=8, before=0, line_spacing=1.333, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing
    if align is not None:
        paragraph.alignment = align


def add_inline_markdown(paragraph, text, *, size=11):
    # Keep citation placeholders and math-like text literal; only handle simple bold/italic.
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, size=size, bold=True)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_font(run, size=size, italic=True)
        else:
            run = paragraph.add_run(part)
            set_font(run, size=size)


def add_heading(doc, text, level):
    if level == 1:
        p = doc.add_paragraph()
        style_paragraph(p, before=18, after=10, line_spacing=1.15)
        run = p.add_run(text)
        set_font(run, size=16, bold=True, color="2E74B5")
    elif level == 2:
        p = doc.add_paragraph()
        style_paragraph(p, before=12, after=6, line_spacing=1.15)
        run = p.add_run(text)
        set_font(run, size=13, bold=True, color="2E74B5")
    else:
        p = doc.add_paragraph()
        style_paragraph(p, before=8, after=4, line_spacing=1.15)
        run = p.add_run(text)
        set_font(run, size=12, bold=True, color="1F4D78")


def build_docx(source=SOURCE, output=OUTPUT):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    lines = source.read_text(encoding="utf-8").splitlines()
    in_code = False

    for i, raw in enumerate(lines):
        line = raw.rstrip()
        if not line:
            continue
        if line.startswith("```"):
            in_code = not in_code
            continue

        if i == 0 and line.startswith("# "):
            p = doc.add_paragraph()
            style_paragraph(p, before=0, after=6, line_spacing=1.15, align=WD_ALIGN_PARAGRAPH.LEFT)
            run = p.add_run(line[2:].strip())
            set_font(run, size=20, bold=True, color="0B2545")
            continue

        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), 1)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), 1)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), 2)
        elif line.startswith("- "):
            p = doc.add_paragraph(style=None)
            p.style = doc.styles["Normal"]
            p.paragraph_format.left_indent = Inches(0.375)
            p.paragraph_format.first_line_indent = Inches(-0.194)
            style_paragraph(p, before=0, after=4, line_spacing=1.208)
            run = p.add_run("• ")
            set_font(run, size=11)
            add_inline_markdown(p, line[2:].strip())
        elif in_code:
            p = doc.add_paragraph()
            style_paragraph(p, after=4, line_spacing=1.0)
            run = p.add_run(line)
            set_font(run, name="Courier New", size=9)
        else:
            p = doc.add_paragraph()
            style_paragraph(p, before=0, after=8, line_spacing=1.333, align=WD_ALIGN_PARAGRAPH.JUSTIFY)
            add_inline_markdown(p, line)

    version_match = re.search(r"_v(\d+)", output.stem)
    version_label = f"v{version_match.group(1)}" if version_match else "working draft"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run(f"Working draft {version_label} | author and declaration sections added; repository/archive URLs pending")
    set_font(r, size=9, color="666666")

    doc.save(output)
    print(output)


if __name__ == "__main__":
    if len(sys.argv) == 3:
        build_docx(Path(sys.argv[1]).resolve(), Path(sys.argv[2]).resolve())
    elif len(sys.argv) == 1:
        build_docx()
    else:
        raise SystemExit("usage: build_manuscript_docx.py [source.md output.docx]")
